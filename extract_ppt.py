from pptx import Presentation
import pandas as pd


file_ppt = "nama_file.pptx"

prs = Presentation(file_ppt)

data = []


for slide_number, slide in enumerate(prs.slides, start=1):

    for shape in slide.shapes:

        if hasattr(shape, "text"):

            text = shape.text.strip()

            if text:

                data.append({
                    "slide": slide_number,
                    "text": text
                })


df = pd.DataFrame(data)

df.to_excel(
    "hasil_extract.xlsx",
    index=False
)


print("Ekstraksi selesai")